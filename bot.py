import asyncio
import datetime
import json
import logging
import os
from aiohttp import web
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import requests

from aiogram import Bot, Dispatcher, F, types
from aiogram.filters import Command
from aiogram.types import FSInputFile, ReplyKeyboardMarkup, KeyboardButton

# ==================== НАСТРОЙКИ ====================
BOT_TOKEN = os.environ.get("BOT_TOKEN", "8841026536:ВАШ_ТОКЕН")
BEARER_TOKEN = os.environ.get(
    "BEARER_TOKEN",
    "eyJ0eXAiOiJKV1QiLCJhbGciOiJSUzI1NiJ9.eyJhdWQiOiJhMmI4MzQyOC0yYmY0LTRiMGQtYTRhZi1lZjMzNDBhZmI0MmYiLCJqdGkiOiI1YTFlNDAyZTA3OWQ2ZjE4MWFhMGQzZmUwNjg2Njg3NDI0ZjBmNDZjMmNjOWEyMDYxZmZlM2IyMTFhYjQ0NjA4NzI0NThmYzZiZmY2MGFhMSIsImlhdCI6MTc5MTc4MjA1MC41MzY1NzksIm5iZiI6MTc5MTc4MjA1MC41MzY1ODEsImV4cCI6MTgyMzMxODA1MC41MTE1OTEsInN1YiI6Ijk3MzYwOTY2Iiwic2NvcGVzIjpbXX0.F9Sfvq2q7jpTojbfub7iMWyD3V2AMhNbS1MxLojhP0GXRzaE1NszGtd3xuXjTuzxa-btw-dDp_JTHIPm83u-8tnS6zB-Gxhb5t4XwTdZN5sxEvilj0kuEeOkx2WmzAEn4uU_1pY4Lsd8s6iyQV6sKcYRDd-NNEY9TgQOj8xwWbHSxBwpi6eKiR0fQaQGEbcYtdAjj9WrdVemsgWBJ27rCDlL1GlZJKXOkCrTWtyt5io7MGrN3UZnkvLXuzGXfmefzoba4PZJ9Z_NRE4NUfPjYxOY0WMz3_McbWk49Nveg85hx0x3fbrG2Cds4WayRu_axUOUJmKgvO9XqZnU0hr9OkWy7wy8wCZ2QybHWVjLp4-ctYgZ5m7RaOIN_5B9X76pr07g3qtXx-PhUeW9SvD2AN13mblOHcDvjS6gKgguhYOfQ8U32MihokCO2anJ3-wWJHH7dSu2A9DVWigUfHuB51ma17er7anbNZTfcAJAA0m023EX0di35pOa5OAGOvz0H7eMLwVYheluG_taI4DeWRPKnLNXr7JTRsGw75UCIC7xKYUQXwNikjChPpqyD01mVOS8W59GkwuGZLgK0rK5sOzZYVmhTkj8gV-3T0k-coDmWmpXYTTptrE7ZgLDJuBg6_DbLux5I_wIjlriqy4fJYLJn2ocv5VYCNUL73qKM5U"
)
SCHOOL_ID = "835"
EDU_YEAR = "2026"

logging.basicConfig(level=logging.INFO)
bot = Bot(token=BOT_TOKEN)
dp = Dispatcher()

# Клавиатура
kb = ReplyKeyboardMarkup(
    keyboard=[
        [KeyboardButton(text="📊 Отчет за сегодня")],
        [KeyboardButton(text="📅 Отчет за вчера")]
    ],
    resize_keyboard=True
)

def fetch_school_statistics(target_date: str):
    url = "https://monitor-api-prod.bilimland.com/api/mpu-v2/school/statistics"
    headers = {
        "accept": "application/json",
        "accept-language": "ru,en;q=0.9,kk;q=0.8",
        "authorization": f"Bearer {BEARER_TOKEN}",
        "origin": "https://www.bilimclass.kz",
        "referer": "https://www.bilimclass.kz/",
        "user-agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)",
        "x-localization": "ru",
        "x-school-id": SCHOOL_ID,
    }
    params = {
        "schoolId": SCHOOL_ID,
        "eduYear": EDU_YEAR,
        "dateStart": target_date,
        "dateEnd": target_date,
        "lang": "ru",
    }
    try:
        response = requests.get(url, headers=headers, params=params, timeout=15)
        response.raise_for_status()
        return response.json()
    except Exception as e:
        logging.error(f"Ошибка запроса: {e}")
        return None

def format_and_save_excel(json_data, target_date: str) -> str:
    teachers_data = json_data.get("data", [])
    if not teachers_data:
        return None

    df = pd.DataFrame(teachers_data)
    column_mapping = {
        "teacherName": "ФИО Учителя",
        "journalIndex": "Индекс ведения (%)",
        "count": "Всего уроков",
        "completed": "Проведено уроков",
        "completedPercent": "% Проведения",
        "themesCount": "Уроков с темой",
        "themesPercent": "% Тем",
        "hwSchedulesCount": "Уроков с ДЗ (всего)",
        "hwWithHomework": "Задано ДЗ",
        "hwWithHomeworkPercentage": "% ДЗ",
        "lessonsToMark": "Уроков для оценок",
        "lessonsMarked": "Оценено уроков",
        "marksPercent": "% Оценок",
    }
    available_cols = [c for c in column_mapping.keys() if c in df.columns]
    df_filtered = df[available_cols].rename(columns=column_mapping)

    file_path = f"BilimClass_Report_{target_date}.xlsx"
    df_filtered.to_excel(file_path, index=False)

    wb = load_workbook(file_path)
    ws = wb.active

    fill_header = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    fill_row_green = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    fill_row_yellow = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    fill_row_red = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

    fill_cell_alert = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    font_cell_alert = Font(name="Calibri", size=10, color="9C0006", bold=True)

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"), right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"), bottom=Side(style="thin", color="D9D9D9")
    )

    ws.row_dimensions[1].height = 28
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    col_journal_idx = headers.index("Индекс ведения (%)") + 1 if "Индекс ведения (%)" in headers else None

    pct_columns = [headers.index(c) + 1 for c in ["% Проведения", "% Тем", "% ДЗ", "% Оценок"] if c in headers]

    for row_idx in range(2, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 22
        journal_val = ws.cell(row=row_idx, column=col_journal_idx).value or 0 if col_journal_idx else 0

        if journal_val >= 100:
            row_fill = fill_row_green
        elif journal_val >= 70:
            row_fill = fill_row_yellow
        else:
            row_fill = fill_row_red

        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.fill = row_fill
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="left" if col_idx == 1 else "center", vertical="center")

        for p_col in pct_columns:
            p_cell = ws.cell(row=row_idx, column=p_col)
            if p_cell.value is not None and p_cell.value < 100:
                p_cell.fill = fill_cell_alert
                p_cell.font = font_cell_alert

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(max_len + 5, 15)

    wb.save(file_path)
    return file_path

@dp.message(Command("start"))
async def cmd_start(message: types.Message):
    await message.answer(
        "👋 Привет! Я бот для генерации отчетов по мониторингу BilimClass.\n"
        "Выберите действие в меню ниже:",
        reply_markup=kb
    )

@dp.message(F.text == "📊 Отчет за сегодня")
async def send_today_report(message: types.Message):
    today_str = datetime.date.today().strftime("%Y-%m-%d")
    await generate_and_send_report(message, today_str)

@dp.message(F.text == "📅 Отчет за вчера")
async def send_yesterday_report(message: types.Message):
    yesterday_str = (datetime.date.today() - datetime.timedelta(days=1)).strftime("%Y-%m-%d")
    await generate_and_send_report(message, yesterday_str)

async def generate_and_send_report(message: types.Message, target_date: str):
    msg = await message.answer(f"⏳ Собираю данные и формирую отчет за {target_date}...")
    
    raw_data = fetch_school_statistics(target_date)
    if not raw_data or not raw_data.get("data"):
        await msg.edit_text(f"❌ Не удалось получить данные за {target_date} (возможно, токен устарел или данных нет).")
        return

    file_path = format_and_save_excel(raw_data, target_date)
    document = FSInputFile(file_path)
    
    await message.answer_document(
        document=document,
        caption=f"✅ Отчет по мониторингу BilimClass за **{target_date}** готов!"
    )
    await msg.delete()

# Обработчик проверки веб-порта для Render
async def handle_ping(request):
    return web.Response(text="Bot is alive!")

async def main():
    app = web.Application()
    app.router.add_get("/", handle_ping)
    runner = web.AppRunner(app)
    await runner.setup()
    port = int(os.environ.get("PORT", 8080))
    site = web.TCPSite(runner, "0.0.0.0", port)
    await site.start()

    await dp.start_polling(bot)

if __name__ == "__main__":
    asyncio.run(main())
